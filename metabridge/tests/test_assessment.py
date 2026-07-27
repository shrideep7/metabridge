"""Migration Assessment Engine: deterministic, parse-only, exports."""
import json
import sys
from pathlib import Path

import pytest

from metabridge.assessment.engine import ASSUMPTIONS, assess
from metabridge.assessment.exports import export_all

ROOT = Path(__file__).resolve().parent.parent / "examples"


@pytest.fixture(autouse=True)
def _iso(tmp_path, monkeypatch):
    monkeypatch.setenv("METABRIDGE_DATA_DIR", str(tmp_path / "iso"))


@pytest.fixture(scope="module")
def pc():
    return assess(str(ROOT / "powercenter_testdata"))


def test_all_sections_present(pc):
    for section in ("executive_summary", "application_inventory",
                    "data_estate_inventory", "object_inventory",
                    "automation_potential", "migration_complexity",
                    "technical_debt", "manual_review_estimate",
                    "resource_estimation", "timeline_estimation",
                    "cost_estimation", "cloud_cost_comparison",
                    "business_impact", "critical_dependencies",
                    "unsupported_features", "migration_risks"):
        assert section in pc, section
    assert "no conversion performed" in pc["determinism_note"]


def test_no_conversion_artifacts(tmp_path, monkeypatch):
    """Parse-only: an assessment must not write conversion output."""
    monkeypatch.setenv("METABRIDGE_DATA_DIR", str(tmp_path / "iso2"))
    before = set((ROOT / "powercenter_testdata").rglob("*"))
    assess(str(ROOT / "powercenter_testdata"))
    assert set((ROOT / "powercenter_testdata").rglob("*")) == before


def test_deterministic():
    a = assess(str(ROOT / "dbt_retail"))
    b = assess(str(ROOT / "dbt_retail"))
    assert json.dumps(a, sort_keys=True) == json.dumps(b,
                                                       sort_keys=True)


def test_executive_summary_consistency(pc):
    es = pc["executive_summary"]
    assert es["objects_total"] == len(pc["object_inventory"])
    assert es["manual_review_items"] == \
        pc["manual_review_estimate"]["items"]
    assert es["estimated_labor_usd"] == \
        pc["cost_estimation"]["labor_usd"]
    assert es["headline"].startswith("%d object" % es["objects_total"])


def test_effort_and_cost_are_assumption_backed(pc):
    r = pc["resource_estimation"]
    assert r["effort_hours"]["total"] == \
        sum(v for k, v in r["effort_hours"].items() if k != "total")
    assert "engineer-week" in r["basis"]
    assert pc["cost_estimation"]["assumptions"]["note"] == \
        ASSUMPTIONS["note"]
    cc = pc["cloud_cost_comparison"]
    n = pc["executive_summary"]["objects_total"]
    assert cc["targets"]["snowflake"]["run_usd_per_month"] == \
        round(n * ASSUMPTIONS["run_cost_usd_per_object_month"][
            "snowflake"], 0)


def test_unsupported_and_risks(pc):
    assert any(u["code"] == "ORCHESTRATION_TASK"
               for u in pc["unsupported_features"])
    assert any(r["risk"] == "Manual porting queue"
               for r in pc["migration_risks"])
    assert all(r["mitigation"] for r in pc["migration_risks"])


def test_sap_and_etl_sources_assess():
    sap = assess(str(ROOT / "sap_landscape"))
    assert sap["source_format"] == "sap"
    assert sap["business_impact"] is not None
    ssis = assess(str(ROOT / "etl_legacy" / "ssis"), "ssis")
    assert ssis["executive_summary"]["objects_total"] >= 2
    assert ssis["data_estate_inventory"]["workflows"]


def test_exports_are_valid_documents(tmp_path, pc):
    files = export_all(pc, str(tmp_path))
    assert set(files) == {"assessment.json", "assessment.xlsx",
                          "assessment.docx", "assessment.pptx",
                          "assessment.pdf"}
    from openpyxl import load_workbook
    wb = load_workbook(str(tmp_path / "assessment.xlsx"))
    assert {"Executive", "Objects", "Risks",
            "Cloud cost"} <= set(wb.sheetnames)
    ws = wb["Objects"]
    assert ws.max_row == len(pc["object_inventory"]) + 1
    from docx import Document
    doc = Document(str(tmp_path / "assessment.docx"))
    assert "Migration Assessment" in doc.paragraphs[0].text
    from pptx import Presentation
    prs = Presentation(str(tmp_path / "assessment.pptx"))
    assert len(prs.slides._sldIdLst) == 7
    assert (tmp_path / "assessment.pdf").read_bytes()[:5] == b"%PDF-"
    assert json.loads((tmp_path / "assessment.json").read_text())[
        "project"] == pc["project"]


# --- API ---------------------------------------------------------------------

@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("METABRIDGE_DATA_DIR", str(tmp_path))
    saved = {m: sys.modules.pop(m, None) for m in ("web.app", "web.auth",
                                                   "web")}
    from fastapi.testclient import TestClient
    import web.app as webapp
    yield TestClient(webapp.app)
    for m, orig in saved.items():
        if orig is not None:
            sys.modules[m] = orig
        else:
            sys.modules.pop(m, None)


def test_assessment_api(client):
    files = [{"name": f.name, "content": f.read_text(errors="replace")}
             for f in (ROOT / "powercenter_testdata").iterdir()
             if f.is_file()]
    r = client.post("/api/assessment", json={"files": files})
    assert r.status_code == 200
    d = r.json()
    aid = d["assessment_id"]
    assert d["executive_summary"]["objects_total"] == 16
    assert len(d["exports"]) == 5

    g = client.get("/api/assessment/%s" % aid)
    assert g.status_code == 200
    assert g.json()["executive_summary"]["objects_total"] == 16

    for fmt, magic in (("pdf", b"%PDF-"), ("xlsx", b"PK"),
                       ("pptx", b"PK"), ("docx", b"PK")):
        e = client.get("/api/assessment/%s/export?format=%s"
                       % (aid, fmt))
        assert e.status_code == 200
        assert e.content[:len(magic)] == magic
    bad = client.get("/api/assessment/%s/export?format=exe" % aid)
    assert bad.status_code == 422
