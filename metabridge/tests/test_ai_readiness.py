"""Enterprise AI Readiness Assessment: deterministic, parse-only."""
import json
import sys
from pathlib import Path

import pytest

from metabridge.ai_readiness.engine import (AI_ASSUMPTIONS,
                                            assess_ai_readiness)
from metabridge.ai_readiness.exports import export_all

ROOT = Path(__file__).resolve().parent.parent / "examples"

_DIMS = ("metadata_quality", "business_glossary", "lineage",
         "data_quality", "master_data", "security", "access_controls",
         "pii", "freshness", "vectorization_readiness",
         "document_quality", "knowledge_graph_readiness",
         "rag_readiness", "llm_readiness", "agent_readiness")


@pytest.fixture(autouse=True)
def _iso(tmp_path, monkeypatch):
    monkeypatch.setenv("METABRIDGE_DATA_DIR", str(tmp_path / "iso"))


@pytest.fixture(scope="module")
def dbt():
    return assess_ai_readiness(str(ROOT / "dbt_retail"))


# --- structure -----------------------------------------------------------

def test_all_dimensions_and_outputs_present(dbt):
    for k in _DIMS:
        assert k in dbt["dimensions"], k
        d = dbt["dimensions"][k]
        assert 0 <= d["score"] <= 100
        assert d["level"] in ("Advanced", "Ready", "Developing",
                              "Foundational", "Not ready")
    for out in ("ai_readiness_score", "rag_readiness",
                "recommended_vector_database", "embedding_strategy",
                "chunking_strategy", "knowledge_graph_strategy",
                "recommended_llm_architecture",
                "estimated_ai_implementation_cost",
                "executive_ai_roadmap"):
        assert out in dbt, out
    assert "no AI in the scores" in dbt["determinism_note"]


def test_score_is_weighted_blend_in_range(dbt):
    sc = dbt["ai_readiness_score"]
    assert 0 <= sc["score"] <= 100
    assert sc["band"] in ("Advanced", "Ready", "Developing",
                          "Foundational", "Not ready")
    assert len(sc["top_blockers"]) == 5
    # blockers sorted ascending by score
    scores = [b["score"] for b in sc["top_blockers"]]
    assert scores == sorted(scores)


def test_deterministic():
    a = assess_ai_readiness(str(ROOT / "dbt_retail"))
    b = assess_ai_readiness(str(ROOT / "dbt_retail"))
    assert json.dumps(a, sort_keys=True) == json.dumps(b, sort_keys=True)


def test_no_llm_and_no_conversion_artifacts(tmp_path, monkeypatch):
    monkeypatch.setenv("METABRIDGE_DATA_DIR", str(tmp_path / "iso2"))
    before = set((ROOT / "dbt_retail").rglob("*"))
    assess_ai_readiness(str(ROOT / "dbt_retail"))
    assert set((ROOT / "dbt_retail").rglob("*")) == before


# --- composites derive from base dimensions ------------------------------

def test_composites_are_derived():
    a = assess_ai_readiness(str(ROOT / "dbt_retail"))
    d = a["dimensions"]
    exp = round(0.30 * d["vectorization_readiness"]["score"]
                + 0.20 * d["document_quality"]["score"]
                + 0.20 * d["metadata_quality"]["score"]
                + 0.15 * d["freshness"]["score"]
                + 0.15 * d["pii"]["score"])
    assert abs(d["rag_readiness"]["score"] - exp) <= 1
    assert a["rag_readiness"]["score"] == d["rag_readiness"]["score"]


# --- outputs are meaningful ----------------------------------------------

def test_rag_gates(dbt):
    gates = dbt["rag_readiness"]["gates"]
    assert len(gates) == 5
    names = {g["gate"] for g in gates}
    assert "Embeddable content" in names
    assert all(isinstance(g["pass"], bool) for g in gates)


def test_cost_is_assumption_backed(dbt):
    c = dbt["estimated_ai_implementation_cost"]
    assert c["assumptions"]["note"] == AI_ASSUMPTIONS["note"]
    b = c["breakdown"]
    # year-1 total = one-time + annual run, both derived
    assert c["total_year_one_usd"] == round(
        c["one_time_usd"] + c["annual_run_usd"], 0)
    assert c["range_year_one_usd"][0] < c["total_year_one_usd"] \
        < c["range_year_one_usd"][1]
    assert b["engineering_usd"] >= 0


def test_roadmap_is_gated(dbt):
    r = dbt["executive_ai_roadmap"]
    assert r["phases"]
    assert r["total_weeks"] == sum(p["weeks"] for p in r["phases"])
    # every phase declares an entry gate
    assert all(p["entry_gate"] for p in r["phases"])


def test_vector_db_decision_tree():
    # a Snowflake estate (via twin) recommends the native store
    twin = {"nodes": [{"kind": "warehouse", "name": "wh",
                       "technology": "snowflake"}]}
    a = assess_ai_readiness(str(ROOT / "dbt_retail"), twin=twin)
    assert "Snowflake" in a["recommended_vector_database"]["recommended"]
    # a Databricks estate recommends Databricks Vector Search
    twin_db = {"nodes": [{"kind": "warehouse", "name": "lh",
                          "technology": "databricks"}]}
    b = assess_ai_readiness(str(ROOT / "dbt_retail"), twin=twin_db)
    assert "Databricks" in b["recommended_vector_database"]["recommended"]
    # no platform, no strict PII -> managed default
    c = assess_ai_readiness(str(ROOT / "dbt_retail"))
    assert c["recommended_vector_database"]["recommended"]


def test_twin_enrichment_lifts_glossary_and_access():
    plain = assess_ai_readiness(str(ROOT / "dbt_retail"))
    twin = {"nodes": [
        {"kind": "domain", "name": "Sales"},
        {"kind": "domain", "name": "Finance"},
        {"kind": "owner", "name": "data@co"},
        {"kind": "warehouse", "name": "wh", "technology": "snowflake"},
    ]}
    rich = assess_ai_readiness(str(ROOT / "dbt_retail"), twin=twin)
    assert rich["dimensions"]["business_glossary"]["score"] > \
        plain["dimensions"]["business_glossary"]["score"]
    assert rich["dimensions"]["access_controls"]["score"] > \
        plain["dimensions"]["access_controls"]["score"]


def test_embedding_strategy_handles_pii():
    sap = assess_ai_readiness(str(ROOT / "sap_landscape"))
    emb = sap["embedding_strategy"]
    assert emb["dimensions"] in (1024, 3072)
    # the strategy must always speak to PII handling before embedding
    assert "redact" in emb["pii_handling"] or "tokenize" in \
        emb["pii_handling"]
    pii_cols = sap["dimensions"]["pii"]["signals"]["pii_columns"]
    # when PII is present without protection evidence, keep it self-hosted
    if pii_cols and sap["dimensions"]["pii"]["signals"]["protected_pct"] \
            == 0:
        assert emb["dimensions"] == 1024
        assert "self-hosted" in emb["model"].lower()


def test_security_and_pii_agree_and_surrogate_keys_dont_inflate():
    """Raw PII reaching a target must read as unprotected on BOTH the
    security and pii dimensions; unrelated md5 surrogate keys must not
    credit protection (the review's core contradiction)."""
    from metabridge.ir.model import (Pipeline, Mapping, Transformation,
                                     Port, TransformationType,
                                     LoadStrategy)
    import metabridge.ai_readiness.engine as E
    p = Pipeline(name="fct", source_format="dbt")
    m = Mapping(name="fct", load_strategy=LoadStrategy.FULL)
    m.transformations = [
        Transformation(name="SRC", type=TransformationType.SOURCE,
                       ports=[Port(name="ssn")]),
        Transformation(name="XF", type=TransformationType.EXPRESSION,
                       ports=[Port(name="order_key",
                                   expression="md5(order_id)")]),
        Transformation(name="TGT", type=TransformationType.TARGET,
                       ports=[Port(name="ssn")]),   # raw, unmasked
    ]
    p.mappings = [m]
    s = E._collect(p, None)
    d = E._dimensions(s)
    assert d["security"]["score"] < 70          # not falsely 'Advanced'
    assert d["pii"]["signals"]["protected_pct"] == 0
    # now mask the target port -> both rise together
    m.transformations[2].ports[0].expression = "sha256(ssn)"
    d2 = E._dimensions(E._collect(p, None))
    assert d2["security"]["score"] >= 90
    assert d2["pii"]["signals"]["protected_pct"] == 100


def test_typed_fields_is_measured_not_constant(dbt):
    # dbt_retail has dates/decimals/integers, so specific typing < 100
    # but > 0 — proving it's a measurement, not the old constant 100
    t = dbt["dimensions"]["metadata_quality"]["signals"][
        "specifically_typed_fields_pct"]
    assert 0 < t < 100


def test_roadmap_gate_text_matches_its_boolean(dbt):
    # each gate string states a numeric threshold and 'met'/'NOT met'
    # consistently — no band-name/threshold contradiction
    for p in dbt["executive_ai_roadmap"]["phases"]:
        g = p["entry_gate"]
        if ">=" in g and "(now" in g:
            import re
            thr = int(re.search(r">=\s*(\d+)", g).group(1))
            now = int(re.search(r"now (\d+)", g).group(1))
            if "NOT met" in g or "deferred" in g or "first" in g:
                assert now < thr, g
            elif " met" in g:
                assert now >= thr, g


def test_dbt_model_descriptions_flow_through(tmp_path):
    # a documented dbt model must lift metadata_quality/glossary
    proj = tmp_path / "proj"
    (proj / "models").mkdir(parents=True)
    (proj / "dbt_project.yml").write_text(
        "name: p\nversion: '1.0'\nconfig-version: 2\nprofile: p\n"
        "model-paths: ['models']\n")
    (proj / "models" / "orders.sql").write_text(
        "select 1 as id from {{ source('raw','o') }}")
    (proj / "models" / "schema.yml").write_text(
        "version: 2\nmodels:\n  - name: orders\n    description: "
        "'The canonical orders fact table for revenue reporting'\n"
        "    columns:\n      - {name: id, data_type: integer}\n")
    (proj / "models" / "sources.yml").write_text(
        "version: 2\nsources:\n  - name: raw\n    schema: RAW\n"
        "    tables:\n      - name: o\n        columns:\n"
        "          - {name: id, data_type: integer}\n")
    a = assess_ai_readiness(str(proj), "dbt")
    assert a["dimensions"]["metadata_quality"]["signals"][
        "described_objects_pct"] > 0
    assert a["dimensions"]["business_glossary"]["score"] > 0


def test_llm_guardrails_always_present(dbt):
    arch = dbt["recommended_llm_architecture"]
    assert arch["pattern"]
    assert len(arch["guardrails"]) >= 4
    assert any("PII" in g for g in arch["guardrails"])
    assert any("citation" in g.lower() or "grounding" in g.lower()
               for g in arch["guardrails"])


def test_multiple_formats_assess():
    for proj, fmt in (("powercenter_testdata", "powercenter"),
                      ("sap_landscape", "sap")):
        a = assess_ai_readiness(str(ROOT / proj))
        assert a["source_format"] == fmt
        assert a["ai_readiness_score"]["score"] >= 0


def test_exports_are_valid_documents(tmp_path, dbt):
    files = export_all(dbt, str(tmp_path))
    assert set(files) == {"ai_readiness.json", "ai_readiness.xlsx",
                          "ai_readiness.pdf"}
    from openpyxl import load_workbook
    wb = load_workbook(str(tmp_path / "ai_readiness.xlsx"))
    assert {"Readiness", "Dimensions", "RAG gates", "Roadmap",
            "Cost"} <= set(wb.sheetnames)
    assert wb["Dimensions"].max_row == len(_DIMS) + 1
    assert (tmp_path / "ai_readiness.pdf").read_bytes()[:5] == b"%PDF-"
    assert json.loads((tmp_path / "ai_readiness.json").read_text())[
        "project"] == dbt["project"]


# --- API -----------------------------------------------------------------

@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("METABRIDGE_DATA_DIR", str(tmp_path))
    saved = {m: sys.modules.pop(m, None)
             for m in ("web.app", "web.auth", "web")}
    from fastapi.testclient import TestClient
    import web.app as webapp
    yield TestClient(webapp.app)
    for m, orig in saved.items():
        if orig is not None:
            sys.modules[m] = orig
        else:
            sys.modules.pop(m, None)


def test_ai_readiness_api(client):
    # send RELATIVE paths so the dbt models/ tree is reconstructed
    files = [{"name": str(f.relative_to(ROOT / "dbt_retail")),
              "content": f.read_text(errors="replace")}
             for f in (ROOT / "dbt_retail").rglob("*") if f.is_file()]
    r = client.post("/api/ai-readiness", json={"files": files})
    assert r.status_code == 200
    d = r.json()
    aid = d["assessment_id"]
    assert 0 <= d["ai_readiness_score"]["score"] <= 100
    assert len(d["exports"]) == 3
    assert len(d["dimensions"]) == 15
    # the tree writer preserved structure -> the project actually parsed
    assert d["dimensions"]["lineage"]["score"] > 0
    assert d["dimensions"]["metadata_quality"]["signals"]["fields"] > 0

    g = client.get("/api/ai-readiness/%s" % aid)
    assert g.status_code == 200
    assert g.json()["project"] == d["project"]

    for fmt, magic in (("pdf", b"%PDF-"), ("xlsx", b"PK"),
                       ("json", b"{")):
        e = client.get("/api/ai-readiness/%s/export?format=%s"
                       % (aid, fmt))
        assert e.status_code == 200
        assert e.content[:len(magic)] == magic
    bad = client.get("/api/ai-readiness/%s/export?format=exe" % aid)
    assert bad.status_code == 422


def test_ai_readiness_api_rejects_bad_input_cleanly(client):
    from fastapi.testclient import TestClient  # noqa: F401
    # malformed JSON -> 422, never 500
    r = client.post("/api/ai-readiness", content=b"{bad",
                    headers={"Content-Type": "application/json"})
    assert r.status_code == 422
    # non-object body -> 422 (must not AttributeError to 500)
    assert client.post("/api/ai-readiness", json=[1, 2]).status_code \
        == 422
    # unknown from_job -> 404, and no job left stranded 'running'
    r2 = client.post("/api/ai-readiness",
                     json={"from_job": "doesnotexist99"})
    assert r2.status_code == 404
    jobs = client.get("/api/jobs").json().get("jobs", [])
    assert not any(j.get("status") == "running" for j in jobs)
    # unknown connection -> 404, not a KeyError-500
    assert client.post("/api/ai-readiness",
                       json={"connection_id": "nope"}).status_code == 404


def test_ai_readiness_api_descends_into_folder_upload(client):
    # a folder pick nests everything under one top dir (webkitRelative
    # path "proj/..") — the build must descend so the project parses
    files = [{"name": "proj/%s" % f.relative_to(ROOT / "dbt_retail"),
              "content": f.read_text(errors="replace")}
             for f in (ROOT / "dbt_retail").rglob("*") if f.is_file()]
    r = client.post("/api/ai-readiness", json={"files": files})
    assert r.status_code == 200
    d = r.json()
    assert d["dimensions"]["lineage"]["score"] > 0        # models parsed
    assert d["dimensions"]["metadata_quality"]["signals"]["fields"] > 0
