"""Security & Compliance Intelligence: deterministic control-gap analysis."""
import json
import sys
from pathlib import Path

import pytest

from metabridge.security.engine import (analyze_security, analyze_from_paths,
                                        _scan_secrets, _is_externalized,
                                        scan_text_secrets, redact_secrets)
from metabridge.security.exports import export_all
from metabridge.twin.model import DigitalTwin
from metabridge.ir.model import (Pipeline, Mapping, Transformation, Port,
                                 TransformationType, SourceTable)

ROOT = Path(__file__).resolve().parent.parent / "examples"

_DIMS = ("iam", "rbac", "secrets", "encryption", "key_management",
         "data_classification", "pii", "pci", "phi")
_FRAMEWORKS = ("gdpr", "hipaa", "pci_dss", "sox", "iso27001", "nist_csf")
_OUTPUTS = ("security_score", "compliance_score", "risk_matrix",
            "recommended_controls", "data_masking_plan",
            "tokenization_plan", "encryption_recommendations",
            "audit_evidence")


@pytest.fixture(autouse=True)
def _iso(tmp_path, monkeypatch):
    monkeypatch.setenv("METABRIDGE_DATA_DIR", str(tmp_path / "iso"))


def _sensitive_pipeline(masked=False):
    """email + ssn + card_number flowing raw (or masked) to a target."""
    p = Pipeline(name="cust_load", source_format="dbt")
    m = Mapping(name="cust")
    tgt_ports = [Port(name="email"), Port(name="ssn"),
                 Port(name="card_number")]
    if masked:
        tgt_ports = [Port(name="email", expression="mask(email)"),
                     Port(name="ssn", expression="sha256(ssn)"),
                     Port(name="card_number",
                          expression="tokenize(card_number)")]
    m.transformations = [
        Transformation(name="SRC", type=TransformationType.SOURCE,
                       properties={"table": "raw"},
                       ports=[Port(name="email"), Port(name="ssn"),
                              Port(name="card_number")]),
        Transformation(name="TGT", type=TransformationType.TARGET,
                       properties={"table": "cust"}, ports=tgt_ports),
    ]
    p.mappings = [m]
    p.sources = [SourceTable(name="raw", columns=[
        Port(name="email"), Port(name="ssn"), Port(name="card_number"),
        Port(name="notes")])]
    return p


# --- structure -----------------------------------------------------------

def test_all_analyses_frameworks_outputs_present():
    r = analyze_security(DigitalTwin("e"), [_sensitive_pipeline()])
    for k in _DIMS:
        assert k in r["analyses"], k
        assert 0 <= r["analyses"][k]["score"] <= 100
    for k in _FRAMEWORKS:
        assert k in r["frameworks"], k
        assert 0 <= r["frameworks"][k]["coverage_pct"] <= 100
    for k in _OUTPUTS:
        assert k in r, k
    assert "not a certified" in r["determinism_note"].lower()
    assert "not a certified attestation" in \
        r["audit_evidence"]["disclaimer"].lower()


def test_deterministic():
    a = analyze_security(DigitalTwin("e"), [_sensitive_pipeline()])
    b = analyze_security(DigitalTwin("e"), [_sensitive_pipeline()])
    assert json.dumps(a, sort_keys=True) == json.dumps(b, sort_keys=True)


# --- classification-driven detection -------------------------------------

def test_pii_pci_phi_detected_and_unprotected():
    r = analyze_security(DigitalTwin("e"), [_sensitive_pipeline()])
    an = r["analyses"]
    assert an["pii"]["signals"]["columns"] >= 2       # email, ssn
    assert an["pci"]["signals"]["columns"] >= 1       # card_number
    # no health data present -> HIPAA identifiers stay PII, NOT phantom PHI
    assert an["phi"]["signals"]["columns"] == 0
    # unprotected -> masking + tokenization plans populated
    masks = {m["object"] for m in r["data_masking_plan"]}
    assert "cust.email" in masks and "cust.ssn" in masks
    toks = {t["object"] for t in r["tokenization_plan"]}
    assert "cust.card_number" in toks


def test_secret_in_quoted_key_dict_is_caught():
    # connection metadata is str()'d dicts / JSON where the KEY is
    # quoted ('password': ...) — the scanner must still catch it
    p = Pipeline(name="etl", source_format="dbt")
    p.metadata = {"connections": [{"name": "wh", "host": "prod",
                                   "user": "admin",
                                   "password": "Pr0dP@ss2026"}]}
    r = analyze_security(DigitalTwin("e"), [p])
    assert r["analyses"]["secrets"]["signals"]["exposures"] >= 1
    assert "Pr0dP@ss2026" not in json.dumps(r)


def test_real_password_with_special_chars_not_cleared():
    # strong passwords containing $-runs / ** must NOT be treated as
    # externalized references and dropped
    for pw in ("Summer$Fall2024", "Pa**word9x", "Xk9$mNq"):
        assert _is_externalized(pw) is False
    for ref in ("${DB_PW}", "<your-password>", "changeme", "env('X')"):
        assert _is_externalized(ref) is True


def test_substring_mask_word_is_not_protection():
    # a raw value whose expression merely CONTAINS a mask fragment
    # (e.g. 'unmask') must NOT be reported protected
    p = Pipeline(name="m", source_format="dbt")
    m = Mapping(name="m")
    m.transformations = [
        Transformation(name="SRC", type=TransformationType.SOURCE,
                       properties={"table": "raw"},
                       ports=[Port(name="card_number")]),
        Transformation(name="TGT", type=TransformationType.TARGET,
                       properties={"table": "t"},
                       ports=[Port(name="card_number",
                                   expression="CONCAT(card_number,'_unmasked')")])]
    p.mappings = [m]
    p.sources = [SourceTable(name="raw",
                             columns=[Port(name="card_number")])]
    r = analyze_security(DigitalTwin("e"), [p])
    assert r["analyses"]["pci"]["signals"]["protected_pct"] == 0
    assert r["tokenization_plan"]          # still needs tokenizing


def test_phi_requires_health_context():
    # a lone email (no health data) is PII, not PHI -> no phantom HIPAA
    p = Pipeline(name="m", source_format="dbt")
    m = Mapping(name="m")
    m.transformations = [
        Transformation(name="SRC", type=TransformationType.SOURCE,
                       properties={"table": "raw"},
                       ports=[Port(name="email")]),
        Transformation(name="TGT", type=TransformationType.TARGET,
                       properties={"table": "t"},
                       ports=[Port(name="email")])]
    p.mappings = [m]
    p.sources = [SourceTable(name="raw", columns=[Port(name="email")])]
    r = analyze_security(DigitalTwin("e"), [p])
    assert r["analyses"]["phi"]["signals"]["columns"] == 0
    assert not any("PHI" in x["risk"] for x in r["risk_matrix"])
    # add genuine health data -> the identifiers become PHI
    p.sources[0].columns.append(Port(name="diagnosis"))
    p.mappings[0].transformations[0].ports.append(Port(name="diagnosis"))
    p.mappings[0].transformations[1].ports.append(Port(name="diagnosis"))
    r2 = analyze_security(DigitalTwin("e"), [p])
    assert r2["analyses"]["phi"]["signals"]["columns"] >= 1


def test_encryption_pct_no_double_count():
    # ssn (masked) + salary (raw): true protection is 1/2 = 50%, not
    # inflated by ssn being counted in both PII and PHI buckets
    p = Pipeline(name="m", source_format="dbt")
    m = Mapping(name="m")
    m.transformations = [
        Transformation(name="SRC", type=TransformationType.SOURCE,
                       properties={"table": "raw"},
                       ports=[Port(name="ssn"), Port(name="salary")]),
        Transformation(name="TGT", type=TransformationType.TARGET,
                       properties={"table": "t"},
                       ports=[Port(name="ssn", expression="sha256(ssn)"),
                              Port(name="salary")])]
    p.mappings = [m]
    p.sources = [SourceTable(name="raw", columns=[Port(name="ssn"),
                                                  Port(name="salary")])]
    r = analyze_security(DigitalTwin("e"), [p])
    assert r["analyses"]["encryption"]["signals"]["protected_pct"] == 50


def test_masking_protects_and_lifts_scores():
    raw = analyze_security(DigitalTwin("e"), [_sensitive_pipeline()])
    masked = analyze_security(DigitalTwin("e"),
                              [_sensitive_pipeline(masked=True)])
    assert masked["analyses"]["encryption"]["score"] > \
        raw["analyses"]["encryption"]["score"]
    assert masked["analyses"]["pci"]["score"] > \
        raw["analyses"]["pci"]["score"]
    # masking plan empties out once everything is protected
    assert not masked["data_masking_plan"]
    assert not masked["tokenization_plan"]


# --- secret scanner (the miss a scanner must never make) -----------------

def test_secret_scanner_catches_real_ignores_externalized():
    p = Pipeline(name="load", source_format="dbt")
    m = Mapping(name="m")
    m.transformations = [
        Transformation(name="SRC", type=TransformationType.SOURCE,
                       properties={"sql_override":
                                   "connect password=hunter2secret to db"}),
        Transformation(name="J", type=TransformationType.JOINER,
                       properties={"condition":
                                   "api_key=sk_live_ABCD1234EFGH5678"}),
        Transformation(name="OK", type=TransformationType.EXPRESSION,
                       properties={"connection": "password=${DB_PW}"}),
        Transformation(name="OK2", type=TransformationType.FILTER,
                       properties={"condition":
                                   "password=<your-password>"}),
    ]
    p.mappings = [m]
    p.metadata = {"connections": ["postgres://svc:realpw123@host/db"]}
    found = {f["type"] for f in _scan_secrets([p])}
    assert "password" in found            # hunter2secret -> caught
    assert "api_key" in found
    assert "connection_string_password" in found
    # a value containing 'secret' is NOT auto-cleared
    assert _is_externalized("hunter2secret") is False
    assert _is_externalized("${DB_PW}") is True
    assert _is_externalized("changeme") is True
    # never leak the value itself
    assert all("redacted" in f["evidence"] for f in _scan_secrets([p]))


def test_scan_text_secrets_reports_location_never_the_value():
    # a single blob (a procedure body fetched from a live source system)
    body = ("CREATE PROCEDURE load() AS $$\n"
            "  conn = connect(password='hunter2secret')\n"
            "  key  = 'api_key=sk_live_ABCD1234EFGH5678'\n"
            "$$")
    found = scan_text_secrets(body, "SALES.LOAD_ORDERS")
    types = {f["type"] for f in found}
    assert "password" in types and "api_key" in types
    assert all(f["location"] == "SALES.LOAD_ORDERS" for f in found)
    assert all("redacted" in f["evidence"] for f in found)
    assert "hunter2secret" not in json.dumps(found)
    # externalized references are not leaks
    assert scan_text_secrets("password=${DB_PW}", "x") == []


def test_redact_secrets_replaces_values_keeps_externalized():
    out = redact_secrets("conn = connect(password='hunter2secret')")
    assert "hunter2secret" not in out
    assert "***REDACTED***" in out
    assert out.startswith("conn = connect(")      # only the VALUE is cut
    # an env/vault reference is not a secret and must survive untouched
    assert redact_secrets("password=${DB_PW}") == "password=${DB_PW}"
    assert redact_secrets("") == ""


def test_secret_in_raw_file_content_is_caught():
    # a secret in raw uploaded content (e.g. a SQL comment the parser
    # drops) must still be detected via the raw-text scan
    r = analyze_security(DigitalTwin("e"), [],
                         raw_texts=[{"name": "model.sql",
                                     "text": "-- password=Sup3rSecretPw\n"
                                             "select 1"}])
    assert r["analyses"]["secrets"]["signals"]["exposures"] >= 1
    assert any("Plaintext secrets" in x["risk"] for x in r["risk_matrix"])
    # the value must never appear in the output
    assert "Sup3rSecretPw" not in json.dumps(r)


def test_plaintext_secret_drives_risk_and_score():
    p = _sensitive_pipeline()
    p.mappings[0].transformations[0].properties["sql_override"] = \
        "select * from raw where access_token=abcd1234efgh5678 x"
    r = analyze_security(DigitalTwin("e"), [p])
    assert r["analyses"]["secrets"]["signals"]["exposures"] >= 1
    assert any("Plaintext secrets" in x["risk"] for x in r["risk_matrix"])


# --- compliance framework math -------------------------------------------

def test_framework_coverage_reflects_controls():
    r = analyze_security(DigitalTwin("e"), [_sensitive_pipeline()])
    for k in _FRAMEWORKS:
        fw = r["frameworks"][k]
        controls = fw["controls"]
        met = sum(1 for c in controls if c["status"] == "met")
        part = sum(1 for c in controls if c["status"] == "partial")
        exp = round(100 * (met + 0.5 * part) / len(controls))
        assert abs(fw["coverage_pct"] - exp) <= 1
    # compliance score is the mean of framework coverages
    mean = round(sum(r["compliance_score"]["by_framework"].values())
                 / len(_FRAMEWORKS))
    assert abs(r["compliance_score"]["score"] - mean) <= 1


def test_risk_matrix_severity_ordering():
    r = analyze_security(DigitalTwin("e"), [_sensitive_pipeline()])
    order = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3}
    sevs = [order[x["severity"]] for x in r["risk_matrix"]]
    assert sevs == sorted(sevs)
    # unprotected PAN must surface as Critical
    assert any(x["severity"] == "Critical" and "PAN" in x["risk"]
               for x in r["risk_matrix"])


def test_audit_evidence_pack():
    r = analyze_security(DigitalTwin("e"), [_sensitive_pipeline()])
    ae = r["audit_evidence"]
    assert set(ae["frameworks"]) == set(_FRAMEWORKS)
    for fw in ae["frameworks"].values():
        assert "controls" in fw and "gaps" in fw
        assert 0 <= fw["coverage_pct"] <= 100
    assert ae["evidence_base"]["plaintext_secrets"] == 0


def test_clean_estate_scores_higher():
    # ownership + masked sensitive data -> better posture
    t = DigitalTwin("clean")
    tbl = t.add_node("table", "fct", owner="data@co")
    t.add_node("owner", "data@co")
    weak = analyze_security(DigitalTwin("e"), [_sensitive_pipeline()])
    strong = analyze_security(t, [_sensitive_pipeline(masked=True)])
    assert strong["security_score"]["score"] > \
        weak["security_score"]["score"]


def test_from_paths_smoke_and_exports(tmp_path):
    r = analyze_from_paths(paths=[str(ROOT / "sap_landscape")])
    assert 0 <= r["security_score"]["score"] <= 100
    assert set(r["frameworks"]) == set(_FRAMEWORKS)
    files = export_all(r, str(tmp_path))
    assert set(files) == {"security.json", "security.xlsx",
                          "security.pdf"}
    from openpyxl import load_workbook
    wb = load_workbook(str(tmp_path / "security.xlsx"))
    assert {"Summary", "Dimensions", "Risk matrix", "Masking plan",
            "Audit evidence"} <= set(wb.sheetnames)
    assert (tmp_path / "security.pdf").read_bytes()[:5] == b"%PDF-"


# --- API -----------------------------------------------------------------

@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("METABRIDGE_DATA_DIR", str(tmp_path))
    saved = {m: sys.modules.pop(m, None)
             for m in ("web.app", "web.auth", "web")}
    from fastapi.testclient import TestClient
    import web.app as webapp
    yield TestClient(webapp.app)
    for m, orig in saved.items():
        if orig is not None:
            sys.modules[m] = orig
        else:
            sys.modules.pop(m, None)


def test_security_api(client):
    files = [{"name": "sap/%s" % f.name,
              "content": f.read_text(errors="replace")}
             for f in (ROOT / "sap_landscape").rglob("*") if f.is_file()]
    r = client.post("/api/security", json={"files": files})
    assert r.status_code == 200
    d = r.json()
    sid = d["security_id"]
    assert 0 <= d["security_score"]["score"] <= 100
    assert len(d["exports"]) == 3
    assert set(d["frameworks"]) == set(_FRAMEWORKS)

    assert client.get("/api/security/%s" % sid).status_code == 200
    for fmt, magic in (("pdf", b"%PDF-"), ("xlsx", b"PK"), ("json", b"{")):
        e = client.get("/api/security/%s/export?format=%s" % (sid, fmt))
        assert e.status_code == 200
        assert e.content[:len(magic)] == magic
    assert client.get("/api/security/%s/export?format=exe"
                      % sid).status_code == 422


def test_security_api_bad_input(client):
    assert client.post("/api/security", content=b"{bad",
                       headers={"Content-Type": "application/json"}
                       ).status_code == 422
    assert client.post("/api/security", json=[1, 2]).status_code == 422
    assert client.post("/api/security", json={}).status_code == 422
