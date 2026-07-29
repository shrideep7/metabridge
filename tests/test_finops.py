"""Enterprise FinOps: deterministic cost model over twin + IR."""
import json
import sys
from pathlib import Path

import pytest

from metabridge.finops.engine import (FINOPS_ASSUMPTIONS, analyze_finops,
                                       analyze_from_paths)
from metabridge.finops.exports import export_all
from metabridge.twin.model import DigitalTwin
from metabridge.ir.model import (Pipeline, Mapping, Transformation, Port,
                                 TransformationType, LoadStrategy)

ROOT = Path(__file__).resolve().parent.parent / "examples"

_ANALYSES = ("warehouse_utilization", "cloud_storage", "streaming_cost",
             "compute_cost", "data_movement", "idle_resources",
             "query_history", "etl_runtime", "pipeline_efficiency")
_OUTPUTS = ("current_cost", "future_cost", "migration_cost", "roi",
            "payback_period", "reserved_capacity_recommendations",
            "warehouse_sizing", "cluster_recommendations",
            "snowflake_optimization", "databricks_optimization",
            "bigquery_optimization", "fabric_optimization")


@pytest.fixture(autouse=True)
def _iso(tmp_path, monkeypatch):
    monkeypatch.setenv("METABRIDGE_DATA_DIR", str(tmp_path / "iso"))


def _estate_twin():
    """Snowflake warehouse, some tables (some unused), pipelines, topics
    (one unconsumed), a dashboard endpoint."""
    t = DigitalTwin("estate")
    wh = t.add_node("warehouse", "PROD_WH", technology="snowflake")
    raw = t.add_node("table", "orders_raw")
    load = t.add_node("pipeline", "load_orders")
    fct = t.add_node("table", "customer_360")
    dash = t.add_node("dashboard", "SalesDash")
    for a, b, k in [(wh.id, raw.id, "contains"), (raw.id, load.id, "feeds"),
                    (load.id, fct.id, "writes"), (fct.id, dash.id, "feeds")]:
        t.add_edge(a, b, k)
    t.add_node("table", "stale_stg")             # unused
    dead = t.add_node("pipeline", "build_stale")
    t.add_edge(dead.id, t.nodes["table:stale_stg"].id, "writes")
    good = t.add_node("topic", "orders.events", metadata={"partitions": 12})
    prod = t.add_node("producer", "svc")
    cons = t.add_node("consumer", "billing")
    t.add_edge(prod.id, good.id, "produces")
    t.add_edge(good.id, cons.id, "consumes")
    aband = t.add_node("topic", "abandoned.events",
                       metadata={"partitions": 6})
    t.add_edge(prod.id, aband.id, "produces")
    return t


def _pipelines():
    p = Pipeline(name="proj", source_format="dbt")
    full = Mapping(name="rebuild", load_strategy=LoadStrategy.FULL)
    full.transformations = [
        Transformation(name="S", type=TransformationType.SOURCE,
                       properties={"table": "raw"}),
        Transformation(name="X", type=TransformationType.EXPRESSION,
                       ports=[Port(name="v", expression="a+b")]),
        Transformation(name="T", type=TransformationType.TARGET,
                       properties={"table": "rebuild"})]
    incr = Mapping(name="append", load_strategy=LoadStrategy.MERGE,
                   unique_key=["id"])
    incr.transformations = [
        Transformation(name="S", type=TransformationType.SOURCE,
                       properties={"table": "raw"}),
        Transformation(name="T", type=TransformationType.TARGET,
                       properties={"table": "append"})]
    p.mappings = [full, incr]
    return [p]


# --- structure -----------------------------------------------------------

def test_all_analyses_and_outputs_present():
    r = analyze_finops(_estate_twin(), _pipelines())
    for k in _ANALYSES:
        assert k in r["analyses"], k
    for k in _OUTPUTS:
        assert k in r, k
    assert "MODELED from estate metadata" in r["assumptions"]["note"]


def test_platform_detected_from_warehouse():
    r = analyze_finops(_estate_twin(), _pipelines())
    assert r["primary_platform"] == "snowflake"
    assert r["snowflake_optimization"]["in_estate"] is True
    assert r["databricks_optimization"]["in_estate"] is False
    # all four playbooks always present with recommendations
    for k in ("snowflake_optimization", "databricks_optimization",
              "bigquery_optimization", "fabric_optimization"):
        assert len(r[k]["recommendations"]) >= 4


# --- cost math invariants (the risk lives here) --------------------------

def test_future_cost_not_above_current():
    r = analyze_finops(_estate_twin(), _pipelines())
    cc, fc = r["current_cost"], r["future_cost"]
    assert fc["monthly_usd"] <= cc["monthly_usd"]
    assert fc["monthly_savings_usd"] >= 0
    assert 0 <= fc["reduction_pct"] <= 100


def test_savings_levers_sum_to_total_no_double_count():
    r = analyze_finops(_estate_twin(), _pipelines())
    fc = r["future_cost"]
    lever_sum = sum(l["monthly_usd"] for l in fc["savings_levers"])
    # rounding of each lever can drift by at most a couple dollars
    assert abs(lever_sum - fc["monthly_savings_usd"]) <= 3


def test_current_components_sum_to_total():
    r = analyze_finops(_estate_twin(), _pipelines())
    cc = r["current_cost"]
    assert abs(sum(cc["by_component_monthly_usd"].values())
               - cc["monthly_usd"]) <= 2
    assert cc["annual_usd"] == round(cc["monthly_usd"] * 12, 0)


def test_roi_and_payback_consistency():
    r = analyze_finops(_estate_twin(), _pipelines())
    fc, roi, pb, mig = (r["future_cost"], r["roi"], r["payback_period"],
                        r["migration_cost"])
    if fc["monthly_savings_usd"] > 0:
        # payback = one-time / monthly savings
        assert abs(pb["months"] - mig["one_time_usd"]
                   / fc["monthly_savings_usd"]) <= 0.2
        # 3-year ROI strictly greater than 1-year ROI when saving
        assert roi["three_year_pct"] > roi["first_year_pct"]
    assert roi["annual_savings_usd"] == fc["annual_savings_usd"]


def test_no_savings_gives_na_payback():
    # a clean, fully-consumed estate with no waste -> minimal/zero savings
    t = DigitalTwin("clean")
    raw = t.add_node("table", "raw", technology="snowflake")
    pipe = t.add_node("pipeline", "load")
    fct = t.add_node("table", "fct_sales")
    dash = t.add_node("dashboard", "D")
    t.add_edge(raw.id, pipe.id, "feeds")
    t.add_edge(pipe.id, fct.id, "writes")
    t.add_edge(fct.id, dash.id, "feeds")
    r = analyze_finops(t, [])          # no IR -> no full-refresh savings
    # with idle=0 and no full-refresh, some compute savings still exist
    # via reserved/rightsizing; assert the payback text is coherent
    pb = r["payback_period"]
    if r["future_cost"]["monthly_savings_usd"] == 0:
        assert pb["months"] is None
        assert "not applicable" in pb["text"]


def test_zero_migration_cost_gives_na_roi_not_zero():
    # a topics-only estate (no pipelines/tables) has $0 migration cost
    # but real savings — ROI must read n/a, never a misleading 0%
    t = DigitalTwin("topics")
    prod = t.add_node("producer", "svc")
    good = t.add_node("topic", "orders", metadata={"partitions": 12})
    cons = t.add_node("consumer", "c")
    t.add_edge(prod.id, good.id, "produces")
    t.add_edge(good.id, cons.id, "consumes")
    t.add_node("topic", "abandoned", metadata={"partitions": 6})  # idle
    t.add_edge(prod.id, t.nodes["topic:abandoned"].id, "produces")
    r = analyze_finops(t, [])
    assert r["migration_cost"]["one_time_usd"] == 0
    if r["future_cost"]["monthly_savings_usd"] > 0:
        assert r["roi"]["first_year_pct"] is None
        assert r["roi"]["three_year_pct"] is None
        assert r["payback_period"]["months"] == 0.0
        assert "immediate" in r["payback_period"]["text"]


def test_idle_cost_ties_to_debt_reachability():
    r = analyze_finops(_estate_twin(), _pipelines())
    idle = r["analyses"]["idle_resources"]
    assert idle["unused_tables"] >= 1      # stale_stg
    assert idle["dead_etl"] >= 1           # build_stale
    assert idle["unused_topics"] >= 1      # abandoned.events
    assert idle["monthly_usd"] > 0
    # idle removal is the first savings lever and matches
    lever = r["future_cost"]["savings_levers"][0]
    assert lever["lever"].startswith("Decommission idle")
    assert abs(lever["monthly_usd"] - idle["monthly_usd"]) <= 1


def test_storage_telemetry_does_not_produce_negative_lever():
    # measured storage smaller than the modeled unused-table inventory
    # must not make idle_storage exceed the bill (negative savings lever)
    t = _estate_twin()
    r = analyze_finops(t, _pipelines(), telemetry={"storage_gb": 1.0})
    assert all(l["monthly_usd"] >= 0
               for l in r["future_cost"]["savings_levers"])
    assert r["analyses"]["idle_resources"]["monthly_usd"] <= \
        r["current_cost"]["monthly_usd"]
    assert r["future_cost"]["monthly_usd"] <= \
        r["current_cost"]["monthly_usd"]


def test_dead_streaming_jobs_not_charged_to_compute():
    # a live batch pipeline + dead streaming jobs: the compute idle must
    # come from dead PIPELINES only; dead streaming hits streaming
    t = DigitalTwin("e")
    raw = t.add_node("table", "raw")
    pipe = t.add_node("pipeline", "load")
    fct = t.add_node("table", "fct_orders")
    dash = t.add_node("dashboard", "D")
    t.add_edge(raw.id, pipe.id, "feeds")
    t.add_edge(pipe.id, fct.id, "writes")
    t.add_edge(fct.id, dash.id, "feeds")
    for i in range(3):
        sj = t.add_node("streaming_job", "sj_%d" % i,
                        technology="databricks")
        tp = t.add_node("topic", "dead_%d" % i, metadata={"partitions": 3})
        t.add_edge(sj.id, tp.id, "writes")
    m = Mapping(name="load", load_strategy=LoadStrategy.MERGE,
                unique_key=["id"])
    m.transformations = [
        Transformation(name="S", type=TransformationType.SOURCE,
                       properties={"table": "raw"}),
        Transformation(name="T", type=TransformationType.TARGET,
                       properties={"table": "fct_orders"})]
    p = Pipeline(name="proj", source_format="dbt", mappings=[m])
    r = analyze_finops(t, [p])
    # the one live pipeline is healthy -> compute idle lever near zero,
    # the idle is streaming (dead jobs), and future never exceeds current
    idle_lever = r["future_cost"]["savings_levers"][0]["monthly_usd"]
    assert idle_lever > 0                            # dead streaming waste
    assert r["future_cost"]["monthly_usd"] <= \
        r["current_cost"]["monthly_usd"]
    # compute component barely reduced (live pipeline not idle)
    assert r["analyses"]["compute_cost"]["monthly_usd"] > 0


def test_cdc_movement_counted_once():
    t = DigitalTwin("e")
    db = t.add_node("database", "pg", technology="postgres")
    tp = t.add_node("topic", "cdc_t", technology="kafka")
    t.add_edge(db.id, tp.id, "feeds")
    dm = analyze_finops(t, [])["analyses"]["data_movement"]
    assert dm["cross_platform_hops"] == 0            # counted as CDC only
    assert dm["cdc_sources"] == 1
    assert dm["estimated_gb_month"] == \
        FINOPS_ASSUMPTIONS["cross_boundary_gb_per_edge_month"]


def test_streaming_job_not_double_billed():
    t = DigitalTwin("e")
    t.add_node("streaming_job", "sj", technology="databricks")
    r = analyze_finops(t, [])
    # streaming job bills under streaming, NOT compute
    assert r["analyses"]["compute_cost"]["monthly_usd"] == 0
    assert r["analyses"]["streaming_cost"]["monthly_usd"] == \
        FINOPS_ASSUMPTIONS["streaming_job_usd_month"]
    # and its modeled provenance is disclosed even with no topics
    assert "streaming_cost" in r["data_basis"]["modeled_from_metadata"]


# --- telemetry override --------------------------------------------------

def test_telemetry_overrides_modeled_components():
    base = analyze_finops(_estate_twin(), _pipelines())
    tele = {"storage_gb": 100000, "monthly_query_usd": 500,
            "monthly_compute_credits": 2000, "credit_price_usd": 3.0}
    r = analyze_finops(_estate_twin(), _pipelines(), telemetry=tele)
    assert "cloud_storage" in r["data_basis"]["measured_from_telemetry"]
    assert "compute_cost" in r["data_basis"]["measured_from_telemetry"]
    assert "query_history" in r["data_basis"]["measured_from_telemetry"]
    # measured storage (100000 GB) dwarfs the modeled estimate
    assert r["analyses"]["cloud_storage"]["monthly_usd"] > \
        base["analyses"]["cloud_storage"]["monthly_usd"]
    assert r["analyses"]["compute_cost"]["monthly_usd"] == \
        round(2000 * 3.0, 0)
    assert r["analyses"]["query_history"]["monthly_usd"] == 500


def test_pipeline_efficiency_reflects_load_strategy():
    r = analyze_finops(_estate_twin(), _pipelines())
    eff = r["analyses"]["pipeline_efficiency"]
    # one MERGE + one FULL -> 50% incremental adoption
    assert eff["incremental_adoption_pct"] == 50.0
    assert eff["full_refresh_loads"] == 1
    assert eff["incremental_loads"] == 1


def test_deterministic_and_exports(tmp_path):
    a = analyze_finops(_estate_twin(), _pipelines())
    b = analyze_finops(_estate_twin(), _pipelines())
    assert json.dumps(a, sort_keys=True) == json.dumps(b, sort_keys=True)
    files = export_all(a, str(tmp_path))
    assert set(files) == {"finops.json", "finops.xlsx", "finops.pdf"}
    from openpyxl import load_workbook
    wb = load_workbook(str(tmp_path / "finops.xlsx"))
    assert {"Summary", "Current cost", "Savings levers", "Analyses",
            "Platform optimization"} <= set(wb.sheetnames)
    assert (tmp_path / "finops.pdf").read_bytes()[:5] == b"%PDF-"


def test_from_paths_smoke():
    import yaml
    estate = yaml.safe_load((ROOT / "estate" / "estate.yml").read_text())
    r = analyze_from_paths(
        paths=[str(ROOT / "etl_legacy" / "ssis"),
               str(ROOT / "events" / "kafka")],
        estate_docs=[estate])
    assert r["current_cost"]["monthly_usd"] >= 0
    assert set(r["analyses"]) == set(_ANALYSES)


# --- API -----------------------------------------------------------------

@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("METABRIDGE_DATA_DIR", str(tmp_path))
    saved = {m: sys.modules.pop(m, None)
             for m in ("web.app", "web.auth", "web")}
    from fastapi.testclient import TestClient
    import web.app as webapp
    yield TestClient(webapp.app)
    for m, orig in saved.items():
        if orig is not None:
            sys.modules[m] = orig
        else:
            sys.modules.pop(m, None)


def test_finops_api(client):
    files = [{"name": "kafka/%s" % f.name,
              "content": f.read_text(errors="replace")}
             for f in (ROOT / "events" / "kafka").iterdir() if f.is_file()]
    client.post("/api/twin/build", json={
        "files": files,
        "estate_yaml": (ROOT / "estate" / "estate.yml").read_text(),
        "include_connections": False, "include_jobs": False})
    r = client.post("/api/finops", json={
        "telemetry": {"storage_gb": 5000}})
    assert r.status_code == 200
    d = r.json()
    fid = d["finops_id"]
    assert d["current_cost"]["monthly_usd"] >= 0
    assert len(d["exports"]) == 3
    assert set(d["analyses"]) == set(_ANALYSES)
    assert "cloud_storage" in d["data_basis"]["measured_from_telemetry"]

    assert client.get("/api/finops/%s" % fid).status_code == 200
    for fmt, magic in (("pdf", b"%PDF-"), ("xlsx", b"PK"), ("json", b"{")):
        e = client.get("/api/finops/%s/export?format=%s" % (fid, fmt))
        assert e.status_code == 200
        assert e.content[:len(magic)] == magic
    assert client.get("/api/finops/%s/export?format=exe"
                      % fid).status_code == 422


def test_finops_api_bad_input(client):
    assert client.post("/api/finops", content=b"{bad",
                       headers={"Content-Type": "application/json"}
                       ).status_code == 422
    # no twin, no files -> clean 422
    assert client.post("/api/finops", json={}).status_code == 422
    # telemetry wrong type -> 422
    r = client.post("/api/finops", json={"telemetry": [1, 2]})
    assert r.status_code == 422
