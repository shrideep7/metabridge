"""Assessment exports — PDF, PowerPoint, Excel, Word, JSON.

All documents render the SAME deterministic assessment dict; nothing is
computed here. openpyxl / python-docx / python-pptx / reportlab are the
canonical pure-Python document libraries.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import List

NAVY = "101D33"
BLUE = "2F6FB4"


# ---------------------------------------------------------------------------
# Excel inventory
# ---------------------------------------------------------------------------

def export_xlsx(a: dict, path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=NAVY)

    def sheet(name, headers, rows):
        ws = wb.create_sheet(name) if wb.sheetnames != ["Sheet"] \
            else wb.active
        ws.title = name
        ws.append(headers)
        for c in ws[1]:
            c.font = head
            c.fill = fill
        for r in rows:
            ws.append(r)
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = \
                max(14, len(str(h)) + 4)
        return ws

    es = a["executive_summary"]
    sheet("Executive", ["Metric", "Value"], [
        ["Project", a["project"]],
        ["Source", a["source_label"]],
        ["Objects", es["objects_total"]],
        ["Automation potential %", es["automation_potential"]],
        ["Complexity", "%s (%s)" % (es["migration_complexity"],
                                    es["complexity_level"])],
        ["Average confidence", es["average_confidence"]],
        ["Manual review items", es["manual_review_items"]],
        ["Technical debt score", es["technical_debt_score"]],
        ["Estimated weeks", es["estimated_weeks"]],
        ["Estimated labor (USD)", es["estimated_labor_usd"]],
    ])
    sheet("Applications",
          ["Application", "Objects", "Manual items", "Avg complexity"],
          [[x["application"], x["objects"], x["manual_items"],
            x["avg_complexity"]] for x in a["application_inventory"]])
    sheet("Objects",
          ["Object", "Kind", "Load strategy", "Transformations",
           "Complexity", "Level", "Automation %", "Confidence",
           "Manual items", "Status"],
          [[o["object"], o["kind"], o["load_strategy"],
            o["transformations"], o["complexity_score"],
            o["complexity_level"], o["automation_percentage"],
            o["conversion_confidence"], o["manual_items"], o["status"]]
           for o in a["object_inventory"]])
    sheet("Data estate", ["Table", "Schema", "Columns"],
          [[t["table"], t["schema"], t["columns"]]
           for t in a["data_estate_inventory"]["tables"]])
    sheet("Unsupported", ["Rule code", "Severity", "Count", "Example"],
          [[u["code"], u["severity"], u["count"], u["example"][:120]]
           for u in a["unsupported_features"]])
    sheet("Risks", ["Risk", "Level", "Evidence", "Mitigation"],
          [[r["risk"], r["level"], r["evidence"], r["mitigation"]]
           for r in a["migration_risks"]])
    sheet("Cloud cost", ["Target", "Run USD/month",
                         "Annual saving vs legacy USD"],
          [[t, v["run_usd_per_month"],
            v["annual_saving_vs_legacy_usd"]]
           for t, v in a["cloud_cost_comparison"]["targets"].items()])
    wb.save(path)


# ---------------------------------------------------------------------------
# Word executive summary
# ---------------------------------------------------------------------------

def export_docx(a: dict, path: str) -> None:
    from docx import Document
    from docx.shared import Pt
    es = a["executive_summary"]
    doc = Document()
    doc.add_heading("Migration Assessment — %s" % a["project"], 0)
    doc.add_paragraph("Source platform: %s · generated deterministically"
                      " by MetaBridge AI (no conversion performed)"
                      % a["source_label"])
    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph(es["headline"])
    doc.add_heading("Key figures", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    for k, v in (("Objects", es["objects_total"]),
                 ("Automation potential",
                  "%s%%" % es["automation_potential"]),
                 ("Migration complexity",
                  "%s (%s)" % (es["migration_complexity"],
                               es["complexity_level"])),
                 ("Average confidence", es["average_confidence"]),
                 ("Manual review items", es["manual_review_items"]),
                 ("Technical debt score",
                  "%s/100" % es["technical_debt_score"]),
                 ("Estimated timeline",
                  "%d weeks" % es["estimated_weeks"]),
                 ("Estimated labor",
                  "$%s" % format(int(es["estimated_labor_usd"]),
                                 ","))):
        row = table.add_row()
        row.cells[0].text = str(k)
        row.cells[1].text = str(v)
    doc.add_heading("Migration risks", level=1)
    for r in a["migration_risks"]:
        p = doc.add_paragraph()
        run = p.add_run("%s (%s): " % (r["risk"], r["level"]))
        run.bold = True
        p.add_run("%s — %s" % (r["evidence"], r["mitigation"]))
    doc.add_heading("Timeline", level=1)
    for ph in a["timeline_estimation"]["phases"]:
        doc.add_paragraph("%s — %d week(s): %s"
                          % (ph["phase"], ph["weeks"], ph["detail"]),
                          style="List Bullet")
    note = doc.add_paragraph(a["assumptions"]["note"])
    note.runs[0].font.size = Pt(8)
    doc.save(path)


# ---------------------------------------------------------------------------
# Board-ready PowerPoint
# ---------------------------------------------------------------------------

def export_pptx(a: dict, path: str) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    es = a["executive_summary"]
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def slide(title):
        s = prs.slides.add_slide(blank)
        bar = s.shapes.add_textbox(Inches(0.5), Inches(0.3),
                                   Inches(12.3), Inches(0.8))
        p = bar.text_frame.paragraphs[0]
        p.text = title
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = RGBColor.from_string(NAVY)
        return s

    def bullets(s, items, top=1.3, size=16):
        box = s.shapes.add_textbox(Inches(0.6), Inches(top),
                                   Inches(12.1), Inches(5.6))
        tf = box.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = "• " + item
            p.font.size = Pt(size)
        return s

    # 1 title
    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(1), Inches(2.6), Inches(11.3),
                              Inches(2))
    p = tb.text_frame.paragraphs[0]
    p.text = "Migration Assessment — %s" % a["project"]
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor.from_string(NAVY)
    p2 = tb.text_frame.add_paragraph()
    p2.text = ("%s · MetaBridge AI · deterministic, pre-conversion"
               % a["source_label"])
    p2.font.size = Pt(18)
    # 2 executive summary
    bullets(slide("Executive summary"), [
        es["headline"],
        "Automation potential: %s%% across %d objects"
        % (es["automation_potential"], es["objects_total"]),
        "Manual review: %d item(s); technical debt %s/100"
        % (es["manual_review_items"], es["technical_debt_score"]),
        "Timeline: ~%d weeks · Labor: ~$%s"
        % (es["estimated_weeks"],
           format(int(es["estimated_labor_usd"]), ","))])
    # 3 scores as a simple chart
    s = slide("Scores")
    try:
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        cd = CategoryChartData()
        cd.categories = ["Automation %", "Confidence",
                         "Complexity", "Tech debt"]
        cd.add_series("score", (es["automation_potential"],
                                es["average_confidence"],
                                es["migration_complexity"],
                                es["technical_debt_score"]))
        s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                           Inches(1), Inches(1.3), Inches(11),
                           Inches(5.4), cd)
    except Exception:  # noqa: BLE001 — chart is decoration, not data
        bullets(s, ["Automation %s%%" % es["automation_potential"],
                    "Confidence %s" % es["average_confidence"],
                    "Complexity %s" % es["migration_complexity"],
                    "Tech debt %s" % es["technical_debt_score"]])
    # 4 inventory
    bullets(slide("Inventory"), [
        "%d objects in %d application group(s)"
        % (es["objects_total"], len(a["application_inventory"]))] + [
        "%s — %d object(s), %d manual, avg complexity %s"
        % (x["application"], x["objects"], x["manual_items"],
           x["avg_complexity"])
        for x in a["application_inventory"][:8]])
    # 5 risks
    bullets(slide("Risk matrix"), [
        "%s [%s] — %s → %s" % (r["risk"], r["level"], r["evidence"],
                               r["mitigation"])
        for r in a["migration_risks"]])
    # 6 roadmap
    bullets(slide("Roadmap & timeline"), [
        "%s — %d week(s): %s" % (p_["phase"], p_["weeks"], p_["detail"])
        for p_ in a["timeline_estimation"]["phases"]] + [
        "Total: ~%d elapsed weeks with %d engineer(s)"
        % (es["estimated_weeks"],
           a["resource_estimation"]["engineers"])])
    # 7 cost
    cc = a["cloud_cost_comparison"]
    bullets(slide("Cost comparison"), [
        "Labor: ~$%s (%s)" % (format(int(es["estimated_labor_usd"]),
                                     ","),
                              "deterministic effort model"),
        "Legacy run rate: $%s/month"
        % format(int(cc["legacy_run_usd_per_month"]), ",")] + [
        "%s: $%s/month (saves $%s/yr vs legacy)"
        % (t, format(int(v["run_usd_per_month"]), ","),
           format(int(v["annual_saving_vs_legacy_usd"]), ","))
        for t, v in list(cc["targets"].items())[:5]] + [
        a["assumptions"]["note"]])
    prs.save(path)


# ---------------------------------------------------------------------------
# PDF assessment report
# ---------------------------------------------------------------------------

def export_pdf(a: dict, path: str) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )
    styles = getSampleStyleSheet()
    h1, h2, body = styles["Title"], styles["Heading2"], styles["BodyText"]
    story = [Paragraph("Migration Assessment — %s" % a["project"], h1),
             Paragraph("Source: %s · MetaBridge AI · deterministic, "
                       "pre-conversion" % a["source_label"], body),
             Spacer(1, 12)]
    es = a["executive_summary"]
    story += [Paragraph("Executive summary", h2),
              Paragraph(es["headline"], body), Spacer(1, 8)]

    def table(headers: List[str], rows: List[list], widths=None):
        data = [headers] + [[Paragraph(str(c), body) for c in r]
                            for r in rows]
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0),
             colors.HexColor("#" + NAVY)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.4,
             colors.HexColor("#c9d2df")),
            ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        return t

    story += [table(["Metric", "Value"], [
        ["Objects", es["objects_total"]],
        ["Automation potential", "%s%%" % es["automation_potential"]],
        ["Complexity", "%s (%s)" % (es["migration_complexity"],
                                    es["complexity_level"])],
        ["Average confidence", es["average_confidence"]],
        ["Manual review items", es["manual_review_items"]],
        ["Technical debt", "%s/100" % es["technical_debt_score"]],
        ["Timeline", "%d weeks" % es["estimated_weeks"]],
        ["Labor", "$%s" % format(int(es["estimated_labor_usd"]), ",")],
    ], widths=[6 * cm, 9 * cm]), Spacer(1, 12)]
    story += [Paragraph("Object inventory", h2),
              table(["Object", "Strategy", "Tx", "Complexity",
                     "Auto %", "Status"],
                    [[o["object"], o["load_strategy"],
                      o["transformations"],
                      "%s (%s)" % (o["complexity_score"],
                                   o["complexity_level"]),
                      o["automation_percentage"], o["status"]]
                     for o in a["object_inventory"][:60]]),
              Spacer(1, 12),
              Paragraph("Unsupported features", h2),
              table(["Rule", "Severity", "Count", "Example"],
                    [[u["code"], u["severity"], u["count"],
                      u["example"][:110]]
                     for u in a["unsupported_features"]] or
                    [["—", "—", 0, "none"]]),
              Spacer(1, 12),
              Paragraph("Migration risks", h2),
              table(["Risk", "Level", "Evidence", "Mitigation"],
                    [[r["risk"], r["level"], r["evidence"],
                      r["mitigation"]] for r in a["migration_risks"]]),
              Spacer(1, 12),
              Paragraph("Timeline", h2),
              table(["Phase", "Weeks", "Detail"],
                    [[p_["phase"], p_["weeks"], p_["detail"]]
                     for p_ in a["timeline_estimation"]["phases"]]),
              Spacer(1, 12),
              Paragraph("Cloud cost comparison (planning figures)", h2),
              table(["Target", "Run $/month", "Annual saving vs legacy"],
                    [[t, v["run_usd_per_month"],
                      v["annual_saving_vs_legacy_usd"]]
                     for t, v in
                     a["cloud_cost_comparison"]["targets"].items()]),
              Spacer(1, 8),
              Paragraph(a["assumptions"]["note"], body)]
    SimpleDocTemplate(path, pagesize=A4,
                      title="MetaBridge Migration Assessment").build(
        story)


EXPORTERS = {"xlsx": export_xlsx, "docx": export_docx,
             "pptx": export_pptx, "pdf": export_pdf}
MEDIA = {"xlsx": "application/vnd.openxmlformats-officedocument."
                 "spreadsheetml.sheet",
         "docx": "application/vnd.openxmlformats-officedocument."
                 "wordprocessingml.document",
         "pptx": "application/vnd.openxmlformats-officedocument."
                 "presentationml.presentation",
         "pdf": "application/pdf", "json": "application/json"}


def export_all(a: dict, out_dir: str) -> List[str]:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    written = []
    (out / "assessment.json").write_text(json.dumps(a, indent=1), encoding="utf-8")
    written.append("assessment.json")
    for ext, fn in EXPORTERS.items():
        name = "assessment.%s" % ext
        fn(a, str(out / name))
        written.append(name)
    return written
