"""Technical Debt exports — JSON, Excel, PDF.

All render the SAME deterministic debt dict; nothing is computed here.
The Excel workbook is the working artifact — engineers execute the
cleanup from the per-category object lists.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import List

NAVY = "101D33"
RUST = "B0422E"

_CATS = ("orphan_datasets", "unused_tables", "unused_columns", "dead_etl",
         "unused_kafka_topics", "unused_dashboards", "unused_apis",
         "unused_process_chains", "duplicate_mappings", "duplicate_sql",
         "duplicate_business_logic", "broken_lineage")


def export_json(d: dict, path: str) -> None:
    Path(path).write_text(json.dumps(d, indent=2))


def export_xlsx(d: dict, path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=NAVY)

    def sheet(name, headers, rows):
        ws = wb.create_sheet(name) if wb.sheetnames != ["Sheet"] \
            else wb.active
        ws.title = name[:31]
        ws.append(headers)
        for c in ws[1]:
            c.font = head
            c.fill = fill
        for r in rows:
            ws.append(r)
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = \
                max(16, min(70, len(str(h)) + 4))
        return ws

    sc = d["technical_debt_score"]
    cost = d["cloud_cost_savings"]
    eff = d["estimated_refactoring_effort"]
    sheet("Summary", ["Metric", "Value"], [
        ["Estate", d["estate"]],
        ["Debt score", "%d / 100 (%s)" % (sc["score"], sc["band"])],
        ["Debt objects", sc["debt_objects"]],
        ["Total estate objects", sc["total_objects"]],
        ["Debt ratio", sc["debt_ratio"]],
        ["Annual cloud saving (USD)", cost["annual_usd"]],
        ["Refactoring effort (hours)", eff["total_hours"]],
        ["Refactoring effort (weeks)", eff["engineer_weeks"]],
        ["Refactoring labor (USD)", eff["labor_usd"]],
    ])
    sheet("By category",
          ["Category", "Count", "Effort (hrs)", "Monthly saving (USD)"],
          [[cat, cnt,
            eff["by_category_hours"].get(cat, 0),
            cost["by_category_monthly_usd"].get(cat, 0)]
           for cat, cnt in sc["by_category"].items()])
    # one row per debt object across all categories
    rows = []
    for cat in _CATS:
        for f in d["findings"].get(cat, []):
            rows.append([f.get("object", ""), f.get("kind", ""),
                         cat.replace("_", " "), f.get("reason", "")])
    sheet("Cleanup items", ["Object", "Kind", "Category", "Why"], rows)
    sheet("Roadmap", ["Phase", "Risk", "Effort (hrs)",
                      "Monthly saving (USD)", "Categories"],
          [[p["phase"], p["risk"], p["effort_hours"],
            p["monthly_savings_usd"],
            "; ".join("%s (%d)" % (i["category"], i["count"])
                      for i in p["items"])]
           for p in d["prioritized_remediation_roadmap"]["phases"]])
    wb.save(path)


def export_pdf(d: dict, path: str) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], fontSize=20,
                        textColor=colors.HexColor("#" + NAVY))
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=13,
                        textColor=colors.HexColor("#" + RUST))
    body = styles["BodyText"]
    small = ParagraphStyle("small", parent=body, fontSize=8,
                           textColor=colors.grey)
    doc = SimpleDocTemplate(path, pagesize=letter,
                            topMargin=0.7 * inch, bottomMargin=0.7 * inch)
    sc = d["technical_debt_score"]
    cost = d["cloud_cost_savings"]
    eff = d["estimated_refactoring_effort"]
    el = [Paragraph("Technical Debt Intelligence", h1),
          Paragraph(d["estate"], small), Spacer(1, 10),
          Paragraph("Debt score: <b>%d / 100</b> — %s"
                    % (sc["score"], sc["band"]), h2),
          Paragraph(sc["headline"], body), Spacer(1, 8)]
    rows = [["Category", "Count"]]
    for cat, cnt in sc["by_category"].items():
        if cnt:
            rows.append([cat, str(cnt)])
    t = Table(rows, colWidths=[3.2 * inch, 1.2 * inch])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#f6f2f3")]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d5dbe4"))]))
    el.append(t)
    el.append(Spacer(1, 12))
    el.append(Paragraph("Cloud saving: <b>$%s/yr</b> &nbsp;·&nbsp; "
                        "effort <b>%s hrs (~%s wk)</b> &nbsp;·&nbsp; "
                        "labor <b>$%s</b>"
                        % (format(int(cost["annual_usd"]), ","),
                           eff["total_hours"], eff["engineer_weeks"],
                           format(int(eff["labor_usd"]), ",")), body))
    el.append(Spacer(1, 10))
    el.append(Paragraph("Prioritized roadmap", h2))
    for p in d["prioritized_remediation_roadmap"]["phases"]:
        el.append(Paragraph(
            "<b>%s</b> [%s] — %s hr, $%s/mo saved: %s"
            % (p["phase"], p["risk"], p["effort_hours"],
               format(int(p["monthly_savings_usd"]), ","),
               ", ".join("%s (%d)" % (i["category"], i["count"])
                         for i in p["items"])), body))
    el.append(Spacer(1, 10))
    el.append(Paragraph(d["coverage_note"], small))
    el.append(Paragraph(d["assumptions"]["note"], small))
    doc.build(el)


MEDIA = {"json": "application/json",
         "xlsx": "application/vnd.openxmlformats-officedocument."
                 "spreadsheetml.sheet",
         "pdf": "application/pdf"}


def export_all(d: dict, out_dir: str) -> List[str]:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    export_json(d, str(out / "tech_debt.json"))
    export_xlsx(d, str(out / "tech_debt.xlsx"))
    export_pdf(d, str(out / "tech_debt.pdf"))
    return ["tech_debt.json", "tech_debt.xlsx", "tech_debt.pdf"]
