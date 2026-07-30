"""AI Readiness exports — JSON, Excel, PDF.

All render the SAME deterministic assessment dict; nothing is computed
here. openpyxl / reportlab are the pure-Python document libraries.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import List

NAVY = "101D33"
TEAL = "0E7C7B"

_DIM_ORDER = [
    "metadata_quality", "business_glossary", "lineage", "data_quality",
    "master_data", "security", "access_controls", "pii", "freshness",
    "vectorization_readiness", "document_quality",
    "knowledge_graph_readiness", "rag_readiness", "llm_readiness",
    "agent_readiness",
]


def _label(k: str) -> str:
    return k.replace("_", " ").title()


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def export_xlsx(a: dict, path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=NAVY)

    def sheet(name, headers, rows):
        ws = wb.create_sheet(name) if wb.sheetnames != ["Sheet"] \
            else wb.active
        ws.title = name
        ws.append(headers)
        for c in ws[1]:
            c.font = head
            c.fill = fill
        for r in rows:
            ws.append(r)
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = \
                max(16, len(str(h)) + 4)
        return ws

    score = a["ai_readiness_score"]
    cost = a["estimated_ai_implementation_cost"]
    sheet("Readiness", ["Metric", "Value"], [
        ["Project", a["project"]],
        ["Source", a["source_label"]],
        ["AI readiness score", "%d / 100 (%s)"
         % (score["score"], score["band"])],
        ["RAG readiness", "%d (%s)" % (a["rag_readiness"]["score"],
                                       a["rag_readiness"]["band"])],
        ["Recommended vector DB",
         a["recommended_vector_database"]["recommended"]],
        ["Embedding model", a["embedding_strategy"]["model"]],
        ["LLM architecture", a["recommended_llm_architecture"]["pattern"]],
        ["Estimated year-1 cost (USD)", cost["total_year_one_usd"]],
        ["Roadmap weeks", a["executive_ai_roadmap"]["total_weeks"]],
    ])
    dims = a["dimensions"]
    sheet("Dimensions", ["Dimension", "Score", "Level", "Weight",
                         "Finding"],
          [[_label(k), dims[k]["score"], dims[k]["level"],
            dims[k]["weight"],
            (dims[k]["findings"][0] if dims[k]["findings"] else "")]
           for k in _DIM_ORDER if k in dims])
    sheet("RAG gates", ["Gate", "Pass", "Detail"],
          [[g["gate"], "YES" if g["pass"] else "no", g["detail"]]
           for g in a["rag_readiness"]["gates"]])
    sheet("Roadmap", ["Phase", "Entry gate", "Weeks", "Deliverables"],
          [[p["phase"], p["entry_gate"], p["weeks"],
            "; ".join(p["deliverables"])]
           for p in a["executive_ai_roadmap"]["phases"]])
    b = cost["breakdown"]
    sheet("Cost", ["Line item", "USD"], [
        ["Initial embedding", b["initial_embedding_usd"]],
        ["Annual re-embedding", b["annual_reembedding_usd"]],
        ["Vector DB / month", b["vector_db_usd_per_month"]],
        ["LLM inference / month", b["llm_inference_usd_per_month"]],
        ["Engineering", b["engineering_usd"]],
        ["One-time total", cost["one_time_usd"]],
        ["Annual run total", cost["annual_run_usd"]],
        ["Year-1 total", cost["total_year_one_usd"]],
        ["Year-1 range low", cost["range_year_one_usd"][0]],
        ["Year-1 range high", cost["range_year_one_usd"][1]],
    ])
    wb.save(path)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def export_pdf(a: dict, path: str) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], fontSize=20,
                        textColor=colors.HexColor("#" + NAVY))
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=13,
                        textColor=colors.HexColor("#" + TEAL))
    body = styles["BodyText"]
    small = ParagraphStyle("small", parent=body, fontSize=8,
                           textColor=colors.grey)

    doc = SimpleDocTemplate(path, pagesize=letter,
                            topMargin=0.7 * inch, bottomMargin=0.7 * inch)
    el: list = []
    score = a["ai_readiness_score"]
    el.append(Paragraph("Enterprise AI Readiness Assessment", h1))
    el.append(Paragraph("%s &nbsp;·&nbsp; %s" % (a["project"],
                                                 a["source_label"]), small))
    el.append(Spacer(1, 10))
    el.append(Paragraph("AI Readiness Score: <b>%d / 100</b> — %s"
                        % (score["score"], score["band"]), h2))
    el.append(Paragraph(score["headline"], body))
    el.append(Spacer(1, 8))

    dims = a["dimensions"]
    rows = [["Dimension", "Score", "Level"]]
    for k in _DIM_ORDER:
        if k in dims:
            rows.append([_label(k), str(dims[k]["score"]),
                         dims[k]["level"]])
    t = Table(rows, colWidths=[2.8 * inch, 0.9 * inch, 1.3 * inch])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#f2f5f9")]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d5dbe4")),
    ]))
    el.append(t)
    el.append(Spacer(1, 12))

    el.append(Paragraph("Recommended architecture", h2))
    vdb = a["recommended_vector_database"]
    arch = a["recommended_llm_architecture"]
    for label, val in [
            ("Vector database", vdb["recommended"]),
            ("Embedding", a["embedding_strategy"]["model"]),
            ("Chunking", a["chunking_strategy"]["primary"]),
            ("Knowledge graph", a["knowledge_graph_strategy"]["store"]),
            ("LLM pattern", arch["pattern"])]:
        el.append(Paragraph("<b>%s:</b> %s" % (label, val), body))
    el.append(Spacer(1, 12))

    el.append(Paragraph("Executive roadmap", h2))
    for p in a["executive_ai_roadmap"]["phases"]:
        el.append(Paragraph("<b>%s</b> (%d wk) — gate: %s"
                            % (p["phase"], p["weeks"], p["entry_gate"]),
                            body))
    cost = a["estimated_ai_implementation_cost"]
    el.append(Spacer(1, 8))
    el.append(Paragraph("Estimated year-1 cost: <b>$%s</b> (range $%s–$%s)"
                        % (format(int(cost["total_year_one_usd"]), ","),
                           format(int(cost["range_year_one_usd"][0]), ","),
                           format(int(cost["range_year_one_usd"][1]), ",")),
                        body))
    el.append(Spacer(1, 10))
    el.append(Paragraph(a["determinism_note"], small))
    el.append(Paragraph(cost["assumptions"]["note"], small))
    doc.build(el)


def export_json(a: dict, path: str) -> None:
    Path(path).write_text(json.dumps(a, indent=2), encoding="utf-8")


MEDIA = {
    "json": "application/json",
    "xlsx": "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def export_all(a: dict, out_dir: str) -> List[str]:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    export_json(a, str(out / "ai_readiness.json"))
    export_xlsx(a, str(out / "ai_readiness.xlsx"))
    export_pdf(a, str(out / "ai_readiness.pdf"))
    return ["ai_readiness.json", "ai_readiness.xlsx", "ai_readiness.pdf"]
