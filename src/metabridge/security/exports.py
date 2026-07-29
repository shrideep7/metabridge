"""Security Intelligence exports — JSON, Excel, PDF.

All render the SAME deterministic security dict; nothing is computed
here. The Excel workbook doubles as the audit-evidence pack.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import List

NAVY = "101D33"
RED = "B0422E"

_FW_LABEL = {"gdpr": "GDPR", "hipaa": "HIPAA", "pci_dss": "PCI-DSS",
             "sox": "SOX", "iso27001": "ISO 27001", "nist_csf": "NIST CSF"}


def export_json(d: dict, path: str) -> None:
    Path(path).write_text(json.dumps(d, indent=2))


def export_xlsx(d: dict, path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=NAVY)

    def sheet(name, headers, rows):
        ws = wb.create_sheet(name) if wb.sheetnames != ["Sheet"] \
            else wb.active
        ws.title = name[:31]
        ws.append(headers)
        for c in ws[1]:
            c.font = head
            c.fill = fill
        for r in rows:
            ws.append(r)
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = \
                max(16, min(80, len(str(h)) + 4))
        return ws

    ss, cs = d["security_score"], d["compliance_score"]
    sheet("Summary", ["Metric", "Value"], [
        ["Estate", d["estate"]],
        ["Security score", "%d / 100 (%s)" % (ss["score"], ss["band"])],
        ["Compliance score", "%d / 100 (%s)" % (cs["score"], cs["band"])],
    ] + [["%s coverage %%" % _FW_LABEL.get(k, k), v]
         for k, v in cs["by_framework"].items()])
    sheet("Dimensions", ["Dimension", "Score", "Level", "Finding"],
          [[k.replace("_", " ").title(), v["score"], v["level"],
            (v["findings"][0] if v["findings"] else "")]
           for k, v in d["analyses"].items()])
    sheet("Risk matrix",
          ["Risk", "Likelihood", "Impact", "Severity", "Affected",
           "Frameworks", "Evidence"],
          [[r["risk"], r["likelihood"], r["impact"], r["severity"],
            r["affected"], "; ".join(r["frameworks"]), r["evidence"]]
           for r in d["risk_matrix"]])
    sheet("Recommended controls",
          ["Priority", "Control", "Frameworks", "Effort"],
          [[c["priority"], c["control"], "; ".join(c["frameworks"]),
            c["effort"]] for c in d["recommended_controls"]])
    sheet("Masking plan", ["Object", "Category", "Technique", "Apply at"],
          [[m["object"], m["category"], m["technique"], m["apply_at"]]
           for m in d["data_masking_plan"]])
    sheet("Tokenization plan", ["Object", "Category", "Technique", "Reason"],
          [[t["object"], t["category"], t["technique"], t["reason"]]
           for t in d["tokenization_plan"]])
    # audit evidence — one row per control across frameworks
    rows = []
    for fw, pack in d["audit_evidence"]["frameworks"].items():
        for c in pack["controls"]:
            rows.append([_FW_LABEL.get(fw, fw), c["ref"], c["control"],
                         c["status"]])
    sheet("Audit evidence",
          ["Framework", "Ref", "Control", "Status"], rows)
    wb.save(path)


def export_pdf(d: dict, path: str) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], fontSize=19,
                        textColor=colors.HexColor("#" + NAVY))
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=13,
                        textColor=colors.HexColor("#" + RED))
    body = styles["BodyText"]
    small = ParagraphStyle("small", parent=body, fontSize=8,
                           textColor=colors.grey)
    doc = SimpleDocTemplate(path, pagesize=letter,
                            topMargin=0.7 * inch, bottomMargin=0.7 * inch)
    ss, cs = d["security_score"], d["compliance_score"]
    el = [Paragraph("Security &amp; Compliance Intelligence", h1),
          Paragraph(d["estate"], small), Spacer(1, 10),
          Paragraph("Security <b>%d/100</b> (%s) &nbsp;·&nbsp; Compliance "
                    "<b>%d/100</b> (%s)" % (ss["score"], ss["band"],
                                            cs["score"], cs["band"]), h2),
          Paragraph(ss["headline"], body), Spacer(1, 8)]
    rows = [["Framework", "Coverage %"]]
    for k, v in cs["by_framework"].items():
        rows.append([_FW_LABEL.get(k, k), str(v)])
    t = Table(rows, colWidths=[3.0 * inch, 1.4 * inch])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d5dbe4"))]))
    el.append(t)
    el.append(Spacer(1, 12))
    el.append(Paragraph("Top risks", h2))
    for r in d["risk_matrix"][:6]:
        el.append(Paragraph("<b>[%s]</b> %s — %s (%s)"
                            % (r["severity"], r["risk"], r["evidence"],
                               ", ".join(r["frameworks"]) or "—"), body))
    el.append(Spacer(1, 10))
    el.append(Paragraph("Recommended controls", h2))
    for c in d["recommended_controls"][:8]:
        el.append(Paragraph("<b>%s</b> %s [%s]"
                            % (c["priority"], c["control"],
                               ", ".join(c["frameworks"])), body))
    if d["data_masking_plan"] or d["tokenization_plan"]:
        el.append(Spacer(1, 8))
        el.append(Paragraph("De-identification: %d column(s) to mask, "
                            "%d to tokenize"
                            % (len(d["data_masking_plan"]),
                               len(d["tokenization_plan"])), body))
    el.append(Spacer(1, 10))
    el.append(Paragraph(d["audit_evidence"]["disclaimer"], small))
    doc.build(el)


MEDIA = {"json": "application/json",
         "xlsx": "application/vnd.openxmlformats-officedocument."
                 "spreadsheetml.sheet",
         "pdf": "application/pdf"}


def export_all(d: dict, out_dir: str) -> List[str]:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    export_json(d, str(out / "security.json"))
    export_xlsx(d, str(out / "security.xlsx"))
    export_pdf(d, str(out / "security.pdf"))
    return ["security.json", "security.xlsx", "security.pdf"]
