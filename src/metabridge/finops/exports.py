"""FinOps exports — JSON, Excel, PDF.

All render the SAME deterministic FinOps dict; nothing is computed here.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import List

NAVY = "101D33"
GREEN = "1E8449"

_PLATFORM_KEYS = [("snowflake_optimization", "Snowflake"),
                  ("databricks_optimization", "Databricks"),
                  ("bigquery_optimization", "BigQuery"),
                  ("fabric_optimization", "Microsoft Fabric")]


def export_json(d: dict, path: str) -> None:
    Path(path).write_text(json.dumps(d, indent=2))


def export_xlsx(d: dict, path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=NAVY)

    def sheet(name, headers, rows):
        ws = wb.create_sheet(name) if wb.sheetnames != ["Sheet"] \
            else wb.active
        ws.title = name[:31]
        ws.append(headers)
        for c in ws[1]:
            c.font = head
            c.fill = fill
        for r in rows:
            ws.append(r)
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = \
                max(16, min(80, len(str(h)) + 4))
        return ws

    cc, fc = d["current_cost"], d["future_cost"]
    roi = d["roi"]
    sheet("Summary", ["Metric", "Value"], [
        ["Estate", d["estate"]],
        ["Primary platform", d["primary_platform"]],
        ["Current cost (USD/mo)", cc["monthly_usd"]],
        ["Current cost (USD/yr)", cc["annual_usd"]],
        ["Future cost (USD/mo)", fc["monthly_usd"]],
        ["Monthly savings (USD)", fc["monthly_savings_usd"]],
        ["Annual savings (USD)", fc["annual_savings_usd"]],
        ["Reduction %", fc["reduction_pct"]],
        ["Migration cost (USD)", d["migration_cost"]["one_time_usd"]],
        ["Payback", d["payback_period"]["text"]],
        ["ROI year-1 %", roi["first_year_pct"]],
        ["ROI 3-year %", roi["three_year_pct"]],
    ])
    sheet("Current cost", ["Component", "USD/month"],
          [[k, v] for k, v in cc["by_component_monthly_usd"].items()])
    sheet("Savings levers", ["Lever", "USD/month"],
          [[l["lever"], l["monthly_usd"]] for l in fc["savings_levers"]])
    an = d["analyses"]
    sheet("Analyses", ["Dimension", "Key metric", "Value"], [
        ["Warehouse utilization", "utilization %",
         an["warehouse_utilization"]["utilization_pct"]],
        ["Cloud storage", "USD/mo", an["cloud_storage"]["monthly_usd"]],
        ["Streaming", "USD/mo", an["streaming_cost"]["monthly_usd"]],
        ["Compute", "USD/mo", an["compute_cost"]["monthly_usd"]],
        ["Data movement", "USD/mo", an["data_movement"]["monthly_usd"]],
        ["Idle resources", "USD/mo", an["idle_resources"]["monthly_usd"]],
        ["Query history", "USD/mo", an["query_history"]["monthly_usd"]],
        ["ETL runtime", "hours/mo",
         an["etl_runtime"]["runtime_hours_month"]],
        ["Pipeline efficiency", "score/100",
         an["pipeline_efficiency"]["efficiency_score"]],
    ])
    rows = []
    for key, label in _PLATFORM_KEYS:
        p = d[key]
        for rec in p["recommendations"]:
            rows.append([label, "yes" if p["in_estate"] else "target",
                         rec])
    sheet("Platform optimization",
          ["Platform", "In estate", "Recommendation"], rows)
    wb.save(path)


def export_pdf(d: dict, path: str) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], fontSize=20,
                        textColor=colors.HexColor("#" + NAVY))
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=13,
                        textColor=colors.HexColor("#" + GREEN))
    body = styles["BodyText"]
    small = ParagraphStyle("small", parent=body, fontSize=8,
                           textColor=colors.grey)
    doc = SimpleDocTemplate(path, pagesize=letter,
                            topMargin=0.7 * inch, bottomMargin=0.7 * inch)
    cc, fc, roi = d["current_cost"], d["future_cost"], d["roi"]
    el = [Paragraph("Enterprise FinOps — Cost &amp; Optimization", h1),
          Paragraph("%s &nbsp;·&nbsp; primary platform: %s"
                    % (d["estate"], d["primary_platform"]), small),
          Spacer(1, 10),
          Paragraph("Current <b>$%s/mo</b> &rarr; optimized <b>$%s/mo</b> "
                    "— save <b>$%s/yr</b> (%s%%)"
                    % (format(int(cc["monthly_usd"]), ","),
                       format(int(fc["monthly_usd"]), ","),
                       format(int(fc["annual_savings_usd"]), ","),
                       fc["reduction_pct"]), h2),
          Paragraph("Migration $%s one-time &nbsp;·&nbsp; payback %s "
                    "&nbsp;·&nbsp; ROI %s%% (yr1) / %s%% (3yr)"
                    % (format(int(d["migration_cost"]["one_time_usd"]), ","),
                       d["payback_period"]["text"],
                       roi["first_year_pct"], roi["three_year_pct"]),
                    body),
          Spacer(1, 10), Paragraph("Cost breakdown &amp; savings", h2)]
    rows = [["Component", "Current $/mo"]]
    for k, v in cc["by_component_monthly_usd"].items():
        rows.append([k, format(int(v), ",")])
    rows.append(["", ""])
    rows.append(["Savings lever", "$/mo"])
    for lv in fc["savings_levers"]:
        rows.append([lv["lever"], format(int(lv["monthly_usd"]), ",")])
    t = Table(rows, colWidths=[3.6 * inch, 1.4 * inch])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d5dbe4"))]))
    el.append(t)
    el.append(Spacer(1, 12))
    el.append(Paragraph("Recommendations", h2))
    rc = d["reserved_capacity_recommendations"]
    el.append(Paragraph("<b>Reserved capacity:</b> %s — save ~$%s/mo (%d%%)"
                        % (rc["mechanism"],
                           format(int(rc["estimated_saving_usd_month"]), ","),
                           rc["estimated_saving_pct"]), body))
    el.append(Paragraph("<b>Warehouse sizing:</b> %s%s"
                        % (d["warehouse_sizing"]["recommended_size"],
                           " (multi-cluster)"
                           if d["warehouse_sizing"]["multi_cluster"]
                           else ""), body))
    for key, label in _PLATFORM_KEYS:
        p = d[key]
        tag = " (in estate)" if p["in_estate"] else " (target-state)"
        el.append(Paragraph("<b>%s%s:</b> %s" % (label, tag,
                            "; ".join(p["recommendations"][:4])), body))
    el.append(Spacer(1, 10))
    el.append(Paragraph(d["determinism_note"], small))
    el.append(Paragraph(d["assumptions"]["note"], small))
    doc.build(el)


MEDIA = {"json": "application/json",
         "xlsx": "application/vnd.openxmlformats-officedocument."
                 "spreadsheetml.sheet",
         "pdf": "application/pdf"}


def export_all(d: dict, out_dir: str) -> List[str]:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    export_json(d, str(out / "finops.json"))
    export_xlsx(d, str(out / "finops.xlsx"))
    export_pdf(d, str(out / "finops.pdf"))
    return ["finops.json", "finops.xlsx", "finops.pdf"]
